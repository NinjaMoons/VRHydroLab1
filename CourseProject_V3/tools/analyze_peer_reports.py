#!/usr/bin/env python3
"""Read-only inventory and text-feature extraction for peer course reports."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SECTION_PATTERNS = {
    "assignment": r"задани[ея] на (курсов|выполн)",
    "annotation": r"\bаннотаци[яи]\b",
    "contents": r"\bсодержание\b|\bоглавление\b",
    "abbreviations": r"список сокращ|обозначени[йя]",
    "introduction": r"\bвведение\b",
    "literature_review": r"обзор.{0,40}(литератур|предмет|науч)|анализ научно",
    "manual_calculation": r"ручн.{0,20}расч|эталонн.{0,20}расч|контрольн.{0,20}расч",
    "geometry": r"геометри",
    "salome": r"\bsalome\b",
    "openfoam": r"\bopenfoam\b",
    "application": r"разработк.{0,25}прилож|программн.{0,20}прилож",
    "results": r"анализ результат|результат.{0,20}расч",
    "comparison": r"сравнен",
    "conclusion": r"\bзаключение\b|\bвыводы\b",
    "references": r"список (использован|литератур)|библиограф",
    "appendices": r"\bприложение [а-яa-z]\b",
}


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def pdf_data(path: Path):
    reader = PdfReader(str(path))
    page_text = [(page.extract_text() or "") for page in reader.pages]
    image_count = 0
    for page in reader.pages:
        try:
            resources = page.get("/Resources", {})
            xobjects = resources.get("/XObject", {})
            xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
            for obj in xobjects.values():
                obj = obj.get_object()
                if obj.get("/Subtype") == "/Image":
                    image_count += 1
        except Exception:
            pass
    return "\n\f\n".join(page_text), len(page_text), image_count, None


def docx_data(path: Path):
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    with zipfile.ZipFile(path) as archive:
        image_count = sum(1 for name in archive.namelist() if name.startswith("word/media/"))
    return "\n".join(parts), None, image_count, len(document.tables)


def report_features(path: Path):
    if path.suffix.lower() == ".pdf":
        text, pages, images, tables = pdf_data(path)
    else:
        text, pages, images, tables = docx_data(path)
    normalized = clean(text)
    lower = normalized.casefold()
    variants = sorted(set(re.findall(r"вариант(?:а|у|ом|е)?\s*(?:№|N)?\s*(\d{1,2})", lower)))
    media = []
    if "вода" in lower or "воды" in lower or "жидкост" in lower:
        media.append("вода/жидкость")
    if "воздух" in lower or re.search(r"\bair\b", lower):
        media.append("воздух")
    stacks = [term for term in (
        "SALOME", "OpenFOAM", "ParaView", "Python", "Node.js", "JavaScript",
        "HTML", "CSS", "Flask", "Django", "FastAPI", "PyQt", "Tkinter",
        "React", "Vue", "C#", "WPF", "MATLAB", "FreeCAD",
    ) if term.casefold() in lower]
    features = {name: bool(re.search(pattern, lower, re.I | re.S)) for name, pattern in SECTION_PATTERNS.items()}
    figure_numbers = [int(value) for value in re.findall(r"рисунок\s+(\d+)", lower)]
    table_numbers = [int(value) for value in re.findall(r"таблиц[аы]\s+(\d+)", lower)]
    listing_numbers = [int(value) for value in re.findall(r"листинг\s+(\d+)", lower)]
    return {
        "path": str(path),
        "format": path.suffix.lower()[1:],
        "size": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "pages": pages,
        "words": len(re.findall(r"[\wА-Яа-яЁё-]+", normalized)),
        "images_embedded_or_xobject": images,
        "tables_docx": tables,
        "max_figure_number": max(figure_numbers, default=0),
        "max_table_number": max(table_numbers, default=0),
        "max_listing_number": max(listing_numbers, default=0),
        "variants": variants,
        "media": media,
        "stack": stacks,
        "sections": features,
        "text": normalized,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--text-dir", type=Path)
    args = parser.parse_args()
    paths = sorted(
        path for path in args.root.rglob("*")
        if path.is_file() and path.suffix.lower() in {".pdf", ".docx"}
        and not path.name.startswith("~$")
    )
    records = []
    for number, path in enumerate(paths, 1):
        print(f"[{number}/{len(paths)}] {path}")
        try:
            record = report_features(path)
        except Exception as error:
            record = {"path": str(path), "error": repr(error)}
        records.append(record)
    if args.text_dir:
        args.text_dir.mkdir(parents=True, exist_ok=True)
        for index, record in enumerate(records, 1):
            if "text" not in record:
                continue
            name = f"{index:02d}_{Path(record['path']).stem}.txt"
            (args.text_dir / name).write_text(record["text"] + "\n", encoding="utf-8")
    args.output.write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"records={len(records)} errors={sum('error' in r for r in records)} output={args.output}")


if __name__ == "__main__":
    main()
