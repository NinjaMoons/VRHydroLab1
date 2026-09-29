#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "docs" / "REPORT_MASTER.md"
DEFAULT_OUT = ROOT / "docs" / "final" / "COURSE_PROJECT_REPORT.docx"


def set_cell_margin(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def font_run(run, name="Times New Roman", size=14, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text, size=14):
    text = text.replace("\\|", "|")
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            font_run(paragraph.add_run(text[pos:match.start()]), size=size)
        token = match.group(0)
        if token.startswith("`"):
            font_run(paragraph.add_run(token[1:-1]), "Consolas", size - 1)
        elif token.startswith("**"):
            font_run(paragraph.add_run(token[2:-2]), size=size, bold=True)
        else:
            font_run(paragraph.add_run(token[1:-1]), size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        font_run(paragraph.add_run(text[pos:]), size=size)


def configure_paragraph(paragraph, *, indent=True, spacing=1.5, after=6, before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    fmt = paragraph.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if spacing == 1.5 else WD_LINE_SPACING.SINGLE
    if spacing != 1.5:
        fmt.line_spacing = spacing
    fmt.first_line_indent = Cm(1.25) if indent else Cm(0)
    fmt.widow_control = True


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    font_run(run, size=12)


def toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Содержание обновляется при открытии документа"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


def toc_begin(paragraph):
    """Begin a Word TOC field whose cached result is written as real paragraphs."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    for node in (begin, instr, separate):
        run._r.append(node)


def toc_end(paragraph):
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_math(document, latex):
    mapping = {
        r"\nabla\cdot\mathbf{U}=0,": "∇·U = 0",
        r"\frac{\partial\mathbf{U}}{\partial t}+\nabla\cdot(\mathbf{U}\mathbf{U})=-\nabla p+\nabla\cdot\left[(\nu+\nu_t)\nabla\mathbf{U}\right],": "∂U/∂t + ∇·(UU) = −∇p + ∇·[(ν + νₜ)∇U]",
        r"\Delta p_{Pa}=\rho\left(\overline p_{in}-\overline p_{out}\right),": "Δp₍Pa₎ = ρ(p̄ᵢₙ − p̄ₒᵤₜ)",
        r"Re=\frac{UL_c}{\nu}.": "Re = U·Lᶜ/ν",
        r"Re_a=\frac{U_{in}a}{\nu},\qquad Re_H=\frac{U_{in}H}{\nu}.": "Reₐ = Uᵢₙ·a/ν,     Reₕ = Uᵢₙ·H/ν",
        r"I=0.16Re_H^{-1/8},\qquad l=0.07H,": "I = 0.16·Reₕ⁻¹ᐟ⁸,     l = 0.07·H",
        r"k=1.5(U_{in}I)^2,\qquad \omega=\frac{\sqrt{k}}{C_\mu^{1/4}l},\quad C_\mu=0.09.": "k = 1.5(UᵢₙI)²,     ω = √k/(Cμ¹ᐟ⁴·l),     Cμ = 0.09",
        r"x_{upper,i}=L-R-(4-i)P,\quad i=1..4,": "xᵤₚₚₑᵣ,ᵢ = L − R − (4 − i)P,     i = 1…4",
        r"x_{lower,i}=x_{upper,i}+S,\quad i=1..3.": "xₗₒwₑᵣ,ᵢ = xᵤₚₚₑᵣ,ᵢ + S,     i = 1…3",
        r"Re_a=1192.8429,\qquad Re_H=2286.2823.": "Reₐ = 1192.8429,     Reₕ = 2286.2823",
    }
    formula = mapping.get(latex.strip(), latex)
    p = document.add_paragraph()
    configure_paragraph(p, indent=False, spacing=1.0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = formula
    mr.append(mt)
    omath.append(mr)
    omath_para.append(omath)
    p._p.append(omath_para)
    return p


def split_row(line):
    line = line.strip()[1:-1]
    cells, current, escaped = [], [], False
    for char in line:
        if escaped:
            current.append(char)
            escaped = False
        elif char == "\\":
            escaped = True
        elif char == "|":
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    cells.append("".join(current).strip())
    return cells


def add_table(document, rows):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for r_index, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        if r_index == 0:
            set_repeat_table_header(row)
        for c_index, value in enumerate(values):
            cell = row.cells[c_index]
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            configure_paragraph(p, indent=False, spacing=1.0, after=0,
                                align=WD_ALIGN_PARAGRAPH.CENTER if r_index == 0 or c_index > 0 else WD_ALIGN_PARAGRAPH.LEFT)
            add_inline(p, value, size=9.5)
            if r_index == 0:
                for run in p.runs:
                    run.bold = True
            tc_pr = cell._tc.get_or_add_tcPr()
            if r_index == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "DDEBF7")
                tc_pr.append(shd)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(document, code_lines):
    for index, line in enumerate(code_lines or [""]):
        p = document.add_paragraph()
        configure_paragraph(p, indent=False, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.right_indent = Cm(0.4)
        p.paragraph_format.keep_with_next = index < len(code_lines) - 1
        p_pr = p._p.get_or_add_pPr()
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "12")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), "5B8790")
        borders.append(left)
        font_run(p.add_run(line if line else " "), "Consolas", 9.5)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    for style_name, size, bold, align in (
        ("Title", 18, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", 16, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 2", 14, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 14, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("Caption", 12, False, WD_ALIGN_PARAGRAPH.CENTER),
    ):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = align
        style.paragraph_format.space_before = Pt(10 if "Heading" in style_name else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    if "Code" not in document.styles:
        document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)


def set_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(1.0)


def add_cover(document):
    p = document.add_paragraph()
    configure_paragraph(p, indent=False, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line, size, bold in (
        ("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ", 12, True),
        ("РОССИЙСКОЙ ФЕДЕРАЦИИ", 12, True),
        ("Федеральное государственное автономное образовательное учреждение", 12, False),
        ("высшего образования", 12, False),
        ("«МОСКОВСКИЙ ПОЛИТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ»", 13, True),
        ("КАФЕДРА «СМАРТ-ТЕХНОЛОГИИ»", 13, True),
    ):
        run = p.add_run(line)
        font_run(run, size=size, bold=bold)
        run.add_break()
    for _ in range(4):
        p.add_run().add_break()
    run = p.add_run("КУРСОВОЙ ПРОЕКТ")
    font_run(run, size=18, bold=True)
    run.add_break()
    run = p.add_run("по дисциплине «Системы инженерного анализа»")
    font_run(run, size=14)
    run.add_break()
    run = p.add_run("по направлению 09.03.01 «Информатика и вычислительная техника»")
    font_run(run, size=12)
    run.add_break()
    run = p.add_run("образовательная программа «Интеграция и программирование в САПР»")
    font_run(run, size=12)
    run.add_break()
    p.add_run().add_break()
    run = p.add_run("Параметрическое приложение для расчёта течения воды в канале\nс шахматным массивом квадратных препятствий")
    font_run(run, size=16, bold=True)
    run.add_break()
    run = p.add_run("Вариант №3")
    font_run(run, size=14, bold=True)
    for _ in range(4):
        run.add_break()

    info = document.add_paragraph()
    configure_paragraph(info, indent=False, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_inline(info, "Выполнил: студент группы 231-324\nШиванко Павел Дмитриевич\n\nПреподаватель:\nЛаврененко Илья Станиславович", size=14)
    info.paragraph_format.left_indent = Cm(8.2)
    for _ in range(4):
        info.add_run().add_break()

    city = document.add_paragraph()
    configure_paragraph(city, indent=False, spacing=1.0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    font_run(city.add_run("Москва, 2026"), size=14)


def add_heading(document, text, level):
    p = document.add_paragraph(style=f"Heading {min(level, 3)}")
    add_inline(p, text, size=16 if level == 1 else 14)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    return p


def build(output, pages, toc_map, source):
    text = source.read_text(encoding="utf-8")
    figure_count = len(re.findall(r"\*\*Рисунок \d+ —", text))
    table_count = len(re.findall(r"\*\*Таблица \d+ —", text))
    listing_count = len(re.findall(r"\*\*Листинг \d+ —", text))
    reference_block = text[text.index("# Список использованных источников"):text.index("# Приложение А")]
    reference_count = len(re.findall(r"(?m)^\d+\. .+$", reference_block))
    appendix_count = len(re.findall(r"(?m)^# Приложение [А-Я] —", text))
    original_lines = text.splitlines()
    toc_entries = []
    for source_line in original_lines:
        match = re.match(r"^(#{1,2})\s+(.*)$", source_line)
        if not match:
            continue
        title = match.group(2).strip()
        if title in {"Пояснительная записка к курсовому проекту", "Титульный лист", "Содержание"}:
            continue
        toc_entries.append((len(match.group(1)), title))
    lines = original_lines
    start = lines.index("## Задание на курсовой проект")
    lines = lines[start:]

    document = Document()
    configure_styles(document)
    set_page(document.sections[0])
    document.core_properties.title = "Курсовой проект — вариант №3"
    document.core_properties.subject = "Системы инженерного анализа"
    document.core_properties.author = "Шиванко Павел Дмитриевич"
    document.core_properties.comments = f"Сформировано из проверенного {source.name}"
    add_cover(document)

    body = document.add_section(WD_SECTION.NEW_PAGE)
    set_page(body)
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    page_field(body.footer.paragraphs[0])
    settings = document.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    forced_pages = {
        "Задание на курсовой проект", "Аннотация", "Содержание", "Список сокращений",
        "Введение", "Описание технического задания", "Заключение",
        "Список использованных источников",
    }
    first_heading = True
    in_abstract = False
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if line.startswith("#"):
            match = re.match(r"^(#{1,3})\s+(.*)$", line)
            level = len(match.group(1))
            title = match.group(2).strip()
            if title == "Титульный лист":
                i += 1
                continue
            is_appendix = title.startswith("Приложение ")
            if not first_heading and (level == 1 or title in forced_pages or is_appendix):
                document.add_page_break()
            first_heading = False
            if title == "Аннотация":
                in_abstract = True
            elif title == "Содержание":
                in_abstract = False
            add_heading(document, title, 1 if level == 1 or title in forced_pages or is_appendix else 2)
            if title == "Аннотация" and pages:
                p = document.add_paragraph()
                configure_paragraph(p, indent=True)
                add_inline(
                    p,
                    f"Объём пояснительной записки: {pages} с.; "
                    f"{figure_count} рис.; {table_count} табл.; "
                    f"{listing_count} листингов; {reference_count} источников; "
                    f"{appendix_count} приложения.",
                )
            if title == "Содержание":
                for entry_index, (entry_level, entry_title) in enumerate(toc_entries):
                    toc = document.add_paragraph()
                    configure_paragraph(toc, indent=False, spacing=1.0, after=0,
                                        align=WD_ALIGN_PARAGRAPH.LEFT)
                    toc.paragraph_format.left_indent = Cm(0.75 if entry_level == 2 and re.match(r"^\d+\.\d+", entry_title) else 0)
                    toc.paragraph_format.tab_stops.add_tab_stop(
                        Cm(15.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
                    )
                    if entry_index == 0:
                        toc_begin(toc)
                    font_run(toc.add_run(entry_title), size=11,
                             bold=entry_level == 1 or not re.match(r"^\d+\.\d+", entry_title))
                    toc.add_run("\t")
                    font_run(toc.add_run(str(toc_map.get(entry_title, "—"))), size=11)
                    if entry_index == len(toc_entries) - 1:
                        toc_end(toc)
                i += 1
                while i < len(lines) and not lines[i].startswith("## Список сокращений"):
                    i += 1
                continue
            i += 1
            continue

        if line == r"\[":
            formula = []
            i += 1
            while i < len(lines) and lines[i].strip() != r"\]":
                formula.append(lines[i].strip())
                i += 1
            add_math(document, "".join(formula))
            i += 1
            continue

        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            add_code(document, code)
            i += 1
            continue

        image_match = re.fullmatch(r"!\[[^\]]*\]\(([^)]+)\)\s*", line)
        if image_match:
            rel = image_match.group(1)
            path = (source.parent / rel).resolve()
            p = document.add_paragraph()
            configure_paragraph(p, indent=False, spacing=1.0, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
            if path.name in {"app_default_report.png", "app_changed_report.png", "app_running_report.png"}:
                width = Cm(10.0)
            elif path.name == "app_complete_u0p15_report.png":
                width = Cm(16.0)
            else:
                width = Cm(16.0)
            p.add_run().add_picture(str(path), width=width)
            p.paragraph_format.keep_with_next = True
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(document, rows)
            continue

        caption = re.fullmatch(r"\*\*((?:Рисунок|Таблица|Листинг)\s+\d+\s+—\s+.*?)\*\*", line)
        if caption:
            p = document.add_paragraph(style="Caption")
            add_inline(p, caption.group(1), size=12)
            p.paragraph_format.keep_with_next = False
            i += 1
            continue

        bullet = re.match(r"^-\s+(.*)$", line)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if bullet or numbered:
            content = (bullet or numbered).group(1 if bullet else 2)
            style = "List Bullet" if bullet else None
            p = document.add_paragraph(style=style)
            configure_paragraph(p, indent=False, spacing=1.5, after=3)
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            add_inline(p, content if bullet else f"{numbered.group(1)}. {content}")
            i += 1
            continue

        p = document.add_paragraph()
        configure_paragraph(p, indent=not line.startswith("**"))
        add_inline(p, line)
        i += 1

    for section in document.sections:
        set_page(section)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--pages", type=int, default=0)
    parser.add_argument("--toc-map", type=Path)
    args = parser.parse_args()
    toc_map = {}
    if args.toc_map:
        toc_map = json.loads(args.toc_map.read_text(encoding="utf-8"))
    build(args.output.resolve(), args.pages, toc_map, args.source.resolve())


if __name__ == "__main__":
    main()
