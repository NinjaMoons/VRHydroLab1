#!/usr/bin/env python3

"""Structural validation for the generated LR2-LR8 DOCX reports."""

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
EXPECTED_IMAGES = {2: 14, 3: 17, 4: 5, 5: 4, 6: 5, 7: 3, 8: 8}
EXPECTED_PAGES = {2: 24, 3: 25, 4: 11, 5: 13, 6: 13, 7: 11, 8: 12}
BAD_MARKERS = ("UNVERIFIED", "[ВСТАВИТЬ", "**", "```", r"\frac", r"\sqrt")


for number in range(2, 9):
    path = ROOT / "final_reports" / f"FINAL_REPORT_LR{number}.docx"
    document = Document(path)
    text = "\n".join(p.text for p in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        zip_error = archive.testzip()
    bad = [marker for marker in BAD_MARKERS if marker in text]
    print(
        f"LR{number}: images={len(document.inline_shapes)}/{EXPECTED_IMAGES[number]}, "
        f"tables={len(document.tables)}, "
        f"short={text.count('Краткий ответ:')}, "
        f"long={text.count('Развёрнутый ответ:')}, "
        f"examples={text.count('Пример из нашей работы:')}, "
        f"oMath={xml.count('<m:oMath')}, bad={bad}, zip_error={zip_error}"
    )

    pdf_path = path.with_suffix(".pdf")
    reader = PdfReader(pdf_path)
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    pdf_bad = [marker for marker in BAD_MARKERS if marker in pdf_text]
    a4_pages = all(
        abs(float(page.mediabox.width) - 595.304) < 1
        and abs(float(page.mediabox.height) - 841.89) < 1
        for page in reader.pages
    )
    print(
        f"     PDF pages={len(reader.pages)}/{EXPECTED_PAGES[number]}, "
        f"A4={a4_pages}, short={pdf_text.count('Краткий ответ:')}, "
        f"long={pdf_text.count('Развёрнутый ответ:')}, "
        f"examples={pdf_text.count('Пример из нашей работы:')}, bad={pdf_bad}"
    )
