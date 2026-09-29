#!/usr/bin/env python3

"""Build formatted DOCX reports LR2-LR8 from verified Markdown sources."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "final_reports"

REPORT_TITLES = {
    2: "Поток в полости при движении верхней грани",
    3: "Многофазный поток",
    4: "Обтекание цилиндра потоком воды",
    5: "Естественная конвекция в полости",
    6: ("Параметрическое построение сопла Лаваля в Salome для анализа "
        "течения газа с использованием инструментария вычислительной "
        "гидродинамики OpenFOAM"),
    7: "Создание веб-приложения для автоматизации расчета сопла Лаваля в OpenFOAM",
    8: "Постобработка результатов в ParaView",
}


def clean_text(text: str) -> str:
    """Normalize characters that are fragile in cross-platform PDF rendering."""
    return (text.replace("\u2011", "-")
                .replace("\u2013", "-")
                .replace("\u2014", "-")
                .replace("\u2212", "-")
                .replace("\u00a0", " "))


def set_run_font(run, name="Times New Roman", size=14, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="D9D9D9", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def add_inline(paragraph, text: str, size=14, default_italic=False):
    text = clean_text(text)
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^]]+\]\([^)]+\))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, italic=default_italic)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, italic=default_italic)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Courier New", size=max(9.5, size - 1), italic=False)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        else:
            label = token[1:token.index("](")]
            run = paragraph.add_run(label)
            set_run_font(run, size=size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, italic=default_italic)


def math_run(text):
    r = OxmlElement("m:r")
    rpr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "i")
    rpr.append(sty)
    r.append(rpr)
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def math_frac(num, den):
    f = OxmlElement("m:f")
    n = OxmlElement("m:num")
    d = OxmlElement("m:den")
    for item in (num if isinstance(num, list) else [math_run(num)]):
        n.append(item)
    for item in (den if isinstance(den, list) else [math_run(den)]):
        d.append(item)
    f.extend((n, d))
    return f


def math_rad(expr):
    rad = OxmlElement("m:rad")
    rad_pr = OxmlElement("m:radPr")
    hide = OxmlElement("m:degHide")
    hide.set(qn("m:val"), "1")
    rad_pr.append(hide)
    # A radical in Office Math requires an explicit degree container even
    # when the degree is hidden.  Without it LibreOffice renders an empty
    # placeholder square instead of the radicand.
    degree = OxmlElement("m:deg")
    element = OxmlElement("m:e")
    for item in expr:
        element.append(item)
    rad.extend((rad_pr, degree, element))
    return rad


def math_sup(base, exponent):
    node = OxmlElement("m:sSup")
    base_node = OxmlElement("m:e")
    sup_node = OxmlElement("m:sup")
    for item in base:
        base_node.append(item)
    for item in exponent:
        sup_node.append(item)
    node.extend((base_node, sup_node))
    return node


def add_equation(document, latex: str):
    latex = latex.strip().rstrip(",")
    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")

    if latex.startswith("Re="):
        items = [math_run("Re = "), math_frac("UD", "ν")]
    elif latex.startswith("p="):
        items = [math_run("p = ρRT")]
    elif latex.startswith("a="):
        items = [math_run("a = "), math_rad([math_run("γRT")]), math_run(",    M = "), math_frac("|U|", "a")]
    elif latex.startswith("M="):
        items = [math_run("M = "), math_frac("|U|", [math_rad([math_run("γRT")])])]
    elif latex.startswith("\\frac{\\partial"):
        items = [
            math_frac("∂α", "∂t"),
            math_run(" + ∇·(αU) + ∇·(α(1-α)Uᵣ) = 0"),
        ]
    elif latex.startswith("\\frac{A}"):
        exponent = [math_frac("γ+1", [math_run("2(γ-1)")])]
        base = [
            math_run("["), math_frac("2", "γ+1"),
            math_run("(1 + "), math_frac("γ-1", "2"),
            math_run("M²)]"),
        ]
        items = [math_frac("A", "A*"), math_run(" = "), math_frac("1", "M"), math_sup(base, exponent)]
    else:
        items = [math_run(clean_text(latex))]

    for item in items:
        omath.append(item)
    omath_para.append(omath)
    para._p.append(omath_para)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(1.25)
    pf.space_after = Pt(0)
    pf.widow_control = True

    for name, size, before, after in (("Title", 16, 0, 0), ("Heading 1", 14, 12, 6), ("Heading 2", 14, 9, 4)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True


def title_paragraph(document, text, size=14, bold=False, space_after=0):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(clean_text(text))
    set_run_font(r, size=size, bold=bold)
    return p


def add_title_page(document, lr):
    title_paragraph(document, "Министерство науки и высшего образования Российской Федерации", 13, True, 6)
    title_paragraph(document, "Федеральное государственное автономное образовательное учреждение", 13)
    title_paragraph(document, "высшего образования", 13, False, 5)
    title_paragraph(document, "«МОСКОВСКИЙ ПОЛИТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ»", 14, True, 2)
    title_paragraph(document, "(МОСКОВСКИЙ ПОЛИТЕХ)", 14, False, 5)
    title_paragraph(document, "КАФЕДРА СМАРТ-ТЕХНОЛОГИИ", 14)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    title_paragraph(document, f"Отчет по выполнению лабораторной работы №{lr}", 15, True, 8)
    title_paragraph(document, f"«{REPORT_TITLES[lr]}»", 15, True, 5)
    title_paragraph(document, "по дисциплине «Системы инженерного анализа»", 14, False, 4)
    title_paragraph(document, "по направлению 09.03.01 Информатика и вычислительная техника", 14, False, 3)
    title_paragraph(document, "Образовательная программа (профиль)", 14)
    title_paragraph(document, "«Интеграция и программирование в САПР»", 14)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(64)

    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(11.5)
    values = [
        ("Преподаватель:", "/____________/ Лаврененко Илья Станиславович /"),
        ("", "      подпись                         ФИО"),
        ("Студент:", "/____________/ Шиванко Павел Дмитриевич, 231-324 /"),
        ("", "      подпись                       ФИО, группа"),
    ]
    for i, (left, right) in enumerate(values):
        for j, value in enumerate((left, right)):
            cell = table.cell(i, j)
            cell.width = table.columns[j].width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 30, 40, 30, 40)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, size=12 if i in (1, 3) else 13, italic=i in (1, 3))
    # Remove all title-page table borders.
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(62)
    title_paragraph(document, "Москва, 2026", 14)
    document.add_page_break()
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ОПИСАНИЕ ВЫПОЛНЕННОЙ РАБОТЫ")
    set_run_font(r, size=15, bold=True)


def add_heading(document, text, level):
    p = document.add_paragraph(style="Heading 1" if level == 2 else "Heading 2")
    p.paragraph_format.first_line_indent = Cm(0)
    if text == "Вопросы для самоконтроля":
        p.paragraph_format.page_break_before = True
    add_inline(p, text, size=14)


def add_body_paragraph(document, text, style=None):
    p = document.add_paragraph(style=style)
    if style in ("List Bullet", "List Number"):
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    add_inline(p, text, size=14)
    return p


def add_markdown_table(document, rows):
    headers = rows[0]
    body = rows[1:]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    total = 16.2
    n = len(headers)
    if n == 2:
        widths = [total * 0.42, total * 0.58]
    elif n == 3:
        widths = [total * 0.28, total * 0.32, total * 0.40]
    elif n == 4:
        widths = [total * 0.27, total * 0.18, total * 0.19, total * 0.36]
    elif n == 5:
        widths = [total * 0.16, total * 0.16, total * 0.18, total * 0.20, total * 0.30]
    else:
        widths = [total / n] * n
    font_size = 10 if n >= 5 else 11
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.width = Cm(widths[col])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, "D9EAF7")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, header, size=font_size)
        for run in p.runs:
            run.bold = True
    for row_index, values in enumerate(body, start=1):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F5F8FA")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if n <= 3 or col < n - 1 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, size=font_size)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def code_caption(code, number):
    lines = [line.strip() for line in code.splitlines() if line.strip()]
    if len(lines) == 1:
        command = lines[0].split()[0]
        return f"Листинг {number} - Команда {command}"
    commands = ", ".join(line.split()[0] for line in lines[:3])
    return f"Листинг {number} - Последовательность команд {commands}"


def add_code_block(document, code, number):
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(3)
    r = caption.add_run(code_caption(code, number))
    set_run_font(r, size=12)

    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6)
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shading)
    for index, line in enumerate(code.splitlines()):
        if index:
            p.add_run().add_break(WD_BREAK.LINE)
        run = p.add_run(clean_text(line))
        set_run_font(run, name="Courier New", size=10.5)


def add_image(document, rel_path, alt):
    path = ROOT / rel_path
    if not path.is_file() or path.stat().st_size == 0:
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        width_px, height_px = image.size
    max_width = 15.8
    max_height = 18.5
    ratio = width_px / max(width_px, height_px)
    width_cm = min(max_width, max_height * width_px / height_px)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def parse_table_line(line):
    content = line.strip().strip("|")
    cells = []
    current = []
    in_code = False
    for char in content:
        if char == "`":
            in_code = not in_code
            current.append(char)
        elif char == "|" and not in_code:
            cells.append(clean_text("".join(current).strip()))
            current = []
        else:
            current.append(char)
    cells.append(clean_text("".join(current).strip()))
    return cells


def build_report(lr):
    source = ROOT / f"FINAL_REPORT_LR{lr}.md"
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])
    document.core_properties.title = f"Лабораторная работа №{lr} {REPORT_TITLES[lr]}"
    document.core_properties.author = "Шиванко Павел Дмитриевич"
    document.core_properties.subject = "Системы инженерного анализа"
    add_page_number(document.sections[0].footer.paragraphs[0])
    add_title_page(document, lr)

    # The title page already contains the Markdown H1 and subtitle.
    start = next(i for i, line in enumerate(lines) if line.startswith("## "))
    i = start
    listing_number = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            add_heading(document, clean_text(line[3:].strip()), 2)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(document, clean_text(line[4:].strip()), 3)
            i += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            add_image(document, image_match.group(2), image_match.group(1))
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            listing_number += 1
            add_code_block(document, "\n".join(code_lines), listing_number)
            i += 1
            continue
        if line == r"\[":
            i += 1
            formula = []
            while i < len(lines) and lines[i].strip() != r"\]":
                formula.append(lines[i].strip())
                i += 1
            add_equation(document, " ".join(formula))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1]):
            table_rows = [parse_table_line(line)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_rows.append(parse_table_line(lines[i]))
                i += 1
            add_markdown_table(document, table_rows)
            continue
        if re.match(r"^- ", line):
            add_body_paragraph(document, line[2:].strip(), "List Bullet")
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            add_body_paragraph(document, re.sub(r"^\d+\.\s+", "", line), "List Number")
            i += 1
            continue
        if re.fullmatch(r"\*Рисунок \d+ .*\*", line):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_together = True
            add_inline(p, line[1:-1], size=12, default_italic=False)
            i += 1
            continue

        # Join ordinary Markdown lines until the next structural element.
        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            candidate = lines[i].rstrip()
            if not candidate:
                break
            if (candidate.startswith(("## ", "### ", "```", "![", "|", "- ")) or
                    candidate == r"\[" or re.match(r"^\d+\. ", candidate) or
                    re.fullmatch(r"\*Рисунок \d+ .*\*", candidate)):
                break
            paragraph_lines.append(candidate)
            i += 1
        add_body_paragraph(document, " ".join(paragraph_lines))

    output = OUTPUT / f"FINAL_REPORT_LR{lr}.docx"
    document.save(output)
    print(output)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    selected = [int(value) for value in sys.argv[1:]] if len(sys.argv) > 1 else list(range(2, 9))
    for lr in selected:
        build_report(lr)


if __name__ == "__main__":
    main()
